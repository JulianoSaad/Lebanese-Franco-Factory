"""Excel round-trip for the review queue: export to .xlsx, import edits back.

The reviewer can export the pending queue, fix the Franco spelling in Excel
(and/or set a decision), then upload the file to record every decision at once.
``openpyxl`` ships in the ``[dashboard]`` extra and is imported lazily.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

HEADERS = ["id", "arabic (source)", "franco (edit here)", "decision", "notes"]
DECISIONS = ("correct", "edit", "reject")


def build_workbook(rows: list[dict[str, Any]]) -> bytes:
    """Return an .xlsx (as bytes) with one editable row per queue item."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "review"

    header_fill = PatternFill("solid", fgColor="0B3D2E")
    header_font = Font(bold=True, color="FFFFFF")
    for col, name in enumerate(HEADERS, 1):
        c = ws.cell(row=1, column=col, value=name)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"

    for i, row in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=str(row.get("id", "")))
        a = ws.cell(row=i, column=2, value=str(row.get("source", "")))
        a.alignment = Alignment(horizontal="right", readingOrder=2)
        ws.cell(row=i, column=3, value=str(row.get("target", "")))
        ws.cell(row=i, column=4, value="")
        ws.cell(row=i, column=5, value=str(row.get("reason", "")))

    if len(rows):
        dv = DataValidation(
            type="list", formula1='"' + ",".join(DECISIONS) + '"', allow_blank=True
        )
        dv.prompt = "correct / edit / reject (blank = skip)"
        dv.promptTitle = "Decision"
        ws.add_data_validation(dv)
        dv.add(f"D2:D{len(rows) + 1}")

    widths = {"A": 30, "B": 34, "C": 40, "D": 14, "E": 30}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    # a second sheet with instructions so the file is self-explanatory
    tips = wb.create_sheet("how-to")
    for i, line in enumerate(
        [
            "Lebanese Franco — offline review",
            "",
            "1. Fix the Franco spelling in column C where needed.",
            "2. Set column D (decision):",
            "     correct = the Franco is right as-is",
            "     edit    = you changed column C",
            "     reject  = drop this phrase",
            "     (blank) = skipped; but if you changed column C it counts as an edit",
            "3. Save the file, then upload it back on the Review page.",
            "",
            "Do not change column A (id) — it links your edit to the phrase.",
        ],
        start=1,
    ):
        tips.cell(row=i, column=1, value=line)
    tips.column_dimensions["A"].width = 70

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_workbook(data: bytes) -> list[dict[str, str]]:
    """Read an uploaded workbook back into ``{id, franco, decision, notes}`` rows."""
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    ws = wb["review"] if "review" in wb.sheetnames else wb.worksheets[0]
    out: list[dict[str, str]] = []
    for r, row in enumerate(ws.iter_rows(values_only=True)):
        if r == 0:  # header
            continue
        if not row or not row[0]:
            continue
        rid = str(row[0]).strip()
        franco = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
        decision = str(row[3]).strip().lower() if len(row) > 3 and row[3] is not None else ""
        notes = str(row[4]).strip() if len(row) > 4 and row[4] is not None else ""
        out.append({"id": rid, "franco": franco, "decision": decision, "notes": notes})
    return out
