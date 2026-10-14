"""Excel export/import round-trip + import decision logic."""

from __future__ import annotations

from io import BytesIO

import pytest

pytest.importorskip("openpyxl")

from lebanese_franco_factory.review import app as review  # noqa: E402
from lebanese_franco_factory.webui import excel  # noqa: E402


def _sample_rows():
    return [
        {"id": "synth_a", "source": "أنا متعب", "target": "ana muta3eb", "reason": "check"},
        {"id": "synth_b", "source": "كيفك؟", "target": "kifak?", "reason": ""},
        {"id": "synth_c", "source": "مرحبا", "target": "mar7aba", "reason": ""},
        {"id": "synth_d", "source": "شكرا", "target": "shukran", "reason": ""},
    ]


def test_build_and_parse_round_trip():
    data = excel.build_workbook(_sample_rows())
    assert data[:2] == b"PK"  # valid xlsx zip
    parsed = excel.parse_workbook(data)
    assert [p["id"] for p in parsed] == ["synth_a", "synth_b", "synth_c", "synth_d"]
    assert parsed[0]["franco"] == "ana muta3eb"


def _edit(data: bytes, edits: dict[str, tuple[str, str]]) -> bytes:
    """Apply {id: (franco, decision)} edits and return new xlsx bytes."""
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(data))
    ws = wb["review"]
    for r in range(2, ws.max_row + 1):
        rid = ws.cell(r, 1).value
        if rid in edits:
            franco, decision = edits[rid]
            ws.cell(r, 3, franco)
            ws.cell(r, 4, decision)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_import_decision_rules(tmp_path, monkeypatch):
    rows = _sample_rows()
    lookup = {r["id"]: {"source": r["source"], "target": r["target"]} for r in rows}
    monkeypatch.setattr(review, "_queue_lookup", lambda: lookup)
    fb = tmp_path / "human_feedback.jsonl"
    monkeypatch.setattr(review, "feedback_path", lambda: fb)

    data = excel.build_workbook(rows)
    # a: changed franco, decision correct  -> should become an EDIT
    # b: unchanged, decision correct       -> CORRECT
    # c: decision reject                   -> REJECT
    # d: changed franco, blank decision    -> auto EDIT
    edited = _edit(
        data,
        {
            "synth_a": ("ana ta3ban", "correct"),
            "synth_b": ("kifak?", "correct"),
            "synth_c": ("mar7aba", "reject"),
            "synth_d": ("shukran ktir", ""),
        },
    )
    summary = review.apply_imported_rows(excel.parse_workbook(edited))
    assert summary == {"correct": 1, "edit": 2, "reject": 1, "skipped": 0, "unknown": 0}

    import json

    recs = {
        json.loads(line)["id"]: json.loads(line)
        for line in fb.read_text(encoding="utf-8").splitlines()
        if line.strip()
    }
    assert recs["synth_a"]["decision"] == "edit"
    assert recs["synth_a"]["corrected"] == "ana ta3ban"
    assert recs["synth_b"]["decision"] == "correct"
    assert recs["synth_c"]["decision"] == "reject"
    assert recs["synth_d"]["decision"] == "edit"
    assert recs["synth_d"]["corrected"] == "shukran ktir"


def test_import_blank_unchanged_is_skipped(tmp_path, monkeypatch):
    rows = _sample_rows()
    lookup = {r["id"]: {"source": r["source"], "target": r["target"]} for r in rows}
    monkeypatch.setattr(review, "_queue_lookup", lambda: lookup)
    monkeypatch.setattr(review, "feedback_path", lambda: tmp_path / "fb.jsonl")
    data = excel.build_workbook(rows)  # no edits, no decisions
    summary = review.apply_imported_rows(excel.parse_workbook(data))
    assert summary["skipped"] == len(rows)
    assert summary["edit"] == 0 and summary["correct"] == 0
